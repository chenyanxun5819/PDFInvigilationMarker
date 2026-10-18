"""生成老师名单Excel模板"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create_teacher_list_template(output_path: Path) -> None:
    """创建老师名单Excel模板"""
    workbook = Workbook()
    
    # 删除默认工作表
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    
    # 创建老师名单工作表 (沿用旧版匯入格式)
    teacher_sheet = workbook.create_sheet("老师名单")
    teacher_headers = ["代号Initial", "教师姓名 Name", "座号", "教职员编号", "组别"]
    
    # 设置表头格式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col_idx, header in enumerate(teacher_headers, start=1):
        cell = teacher_sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # 添加示例数据行
    sample_teacher_data = [
        ["A01", "张老师", "1", "T001", "语文"],
        ["A02", "李老师", "2", "T002", "数学"],
        ["A03", "王老师", "3", "T003", "英文"],
    ]
    
    for row_idx, row_data in enumerate(sample_teacher_data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = teacher_sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 设置列宽
    teacher_sheet.column_dimensions["A"].width = 12
    teacher_sheet.column_dimensions["B"].width = 15
    teacher_sheet.column_dimensions["C"].width = 10
    teacher_sheet.column_dimensions["D"].width = 15
    teacher_sheet.column_dimensions["E"].width = 15
    
    # 保存文件
    workbook.save(output_path)
