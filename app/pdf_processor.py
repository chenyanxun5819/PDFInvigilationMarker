from __future__ import annotations

import re
from collections import defaultdict
from functools import lru_cache
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

try:
    import pymupdf
except ModuleNotFoundError:
    import fitz as pymupdf

from .models import AppConfig, MatchOccurrence, ProcessResult, TeacherMatch, TeacherRecord
from .utils import normalize_seat_no

COMMON_FONT_CANDIDATES = (
    r"C:\Windows\Fonts\msjh.ttc",
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\mingliu.ttc",
    r"C:\Windows\Fonts\simsun.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
)
TILE_ROWS = 4
TILE_COLS = 4
TILE_OVERLAP = 40
TILE_ZOOM_MULTIPLIER = 2.75
TEXT_ROW_Y_TOLERANCE = 3.0
OCR_ROW_Y_TOLERANCE = 6.0
TEXT_NEARBY_ROW_TOLERANCE = 18.0
OCR_NEARBY_ROW_TOLERANCE = 28.0
OCR_EFFECTIVE_CONFIDENCE_CAP = 0.48
CODE_NORMALIZATION_MAP = {
    "粱": "梁",
}
EXCLUDED_SCHEDULE_TOKENS = {
    "FORM", "CLASS", "VENUE", "PAX",
    "MPH", "HOSTELHALL", "HOSTEL", "HALL",
    "B201", "B202", "B203", "B204",
    "A", "B", "C", "D", "E", "F", "G", "H", "I", "J",
    "A3", "B3", "C3", "J3", "S3", "S1", "S2", "J1", "J2",
    "MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN",
    "OF", "THE", "UEC",
}


def process_pdf(
    input_pdf: Path,
    output_pdf: Path,
    config: AppConfig,
    progress_callback=None,
) -> ProcessResult:
    if not input_pdf.exists():
        raise FileNotFoundError(f"找不到 PDF：{input_pdf}")

    output_teachers = resolve_output_teachers(input_pdf, config.teachers, config, progress_callback)
    unresolved_teachers = tuple(teacher for teacher in output_teachers if teacher.unresolved)
    teacher_map = build_teacher_lookup(config.teachers + unresolved_teachers)
    seat_map = {normalize_seat_no(item.seat_no): item.block for item in config.seat_blocks}
    font_path = resolve_font_path(config.label_font_path)
    ocr_engine: object | None = None
    occurrences_by_teacher: dict[str, list[MatchOccurrence]] = defaultdict(list)
    warnings: list[str] = []
    text_detected_codes: set[str] = set()

    with pymupdf.open(input_pdf) as source_doc:
        total_pages = source_doc.page_count
        for page_index in range(total_pages):
            if progress_callback:
                progress_callback(f"识别第 {page_index + 1}/{total_pages} 页...")
            page = source_doc[page_index]
            matches, used_text_layer, used_ocr = detect_page_matches(page, ocr_engine, config, teacher_map)
            if used_ocr and config.text_only_mode and source_doc.is_pdf:
                if progress_callback:
                    progress_callback("仅文字层模式已开启：本页跳过 OCR 回退。")
                matches = {}
                used_ocr = False
            if used_ocr and ocr_engine is None:
                ocr_engine = create_ocr_engine()
                matches, used_text_layer, _ = detect_page_matches(page, ocr_engine, config, teacher_map)
                if progress_callback:
                    progress_callback("本页没有足够文字层，已回退使用 OCR 定位。")
            for teacher_code, item in matches.items():
                occurrences_by_teacher[teacher_code].extend(item)
                if used_text_layer:
                    text_detected_codes.add(teacher_code)

        output_teachers = resolve_final_output_teachers(
            output_teachers,
            teacher_map,
            text_detected_codes,
        )
        if text_detected_codes and progress_callback:
            progress_callback(f"PDF 文字层已直接命中 {len(text_detected_codes)} 位老师，优先采用文字座标定位。")

        matched_items = build_teacher_matches(
            output_teachers,
            occurrences_by_teacher,
            seat_map,
            config,
            warnings,
        )
        if not matched_items:
            raise ValueError("没有从 PDF 中识别到任何可匹配的老师代号。请检查 PDF 清晰度或调整设定档。")

        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        with pymupdf.open() as output_doc:
            for teacher_match in matched_items:
                if progress_callback:
                    progress_callback(
                        f"输出 {teacher_match.teacher.initial} {teacher_match.teacher.name}，"
                        f"共 {len(resolve_target_page_indexes(teacher_match, source_doc.page_count))} 页..."
                    )
                append_teacher_pages(output_doc, source_doc, teacher_match, config, font_path)

            output_doc.save(output_pdf, garbage=4, deflate=True)

    unmatched_codes = tuple(
        item.initial for item in output_teachers if normalize_initial(item.initial) not in occurrences_by_teacher
    )
    unresolved_codes = tuple(item.initial for item in output_teachers if item.unresolved)
    return ProcessResult(
        output_path=output_pdf,
        matched_teachers=tuple(matched_items),
        unmatched_teacher_codes=unmatched_codes,
        unresolved_teacher_codes=unresolved_codes,
        warnings=tuple(warnings),
    )


def resolve_output_teachers(
    input_pdf: Path,
    teachers: tuple[TeacherRecord, ...],
    config: AppConfig,
    progress_callback=None,
) -> tuple[TeacherRecord, ...]:
    include_unresolved = not is_image_input_path(input_pdf)
    pdf_codes = detect_input_pdf_codes(input_pdf, teachers, config)
    if pdf_codes:
        resolved = resolve_teacher_records_from_codes(pdf_codes, teachers, include_unresolved=include_unresolved)
        if progress_callback:
            progress_callback(f"已从输入档案自动去重，共 {len(resolved)} 位老师，将依此输出页数。")
        return resolved

    schedule_codes = load_schedule_teacher_codes(input_pdf.parent, teachers)
    if not schedule_codes:
        return teachers

    resolved = resolve_teacher_records_from_codes(schedule_codes, teachers, include_unresolved=include_unresolved)
    if progress_callback:
        progress_callback(f"已从排班 Excel 自动去重，共 {len(resolved)} 位老师，将依此输出页数。")
    return tuple(resolved)


def resolve_final_output_teachers(
    schedule_teachers: tuple[TeacherRecord, ...],
    teacher_map: dict[str, TeacherRecord],
    text_detected_codes: set[str],
) -> tuple[TeacherRecord, ...]:
    if text_detected_codes:
        ordered_teachers = [
            teacher
            for teacher in schedule_teachers
            if normalize_initial(teacher.initial) in text_detected_codes
        ]
        ordered_codes = {normalize_initial(item.initial) for item in ordered_teachers}
        for code in sorted(text_detected_codes):
            if code not in ordered_codes and code in teacher_map:
                ordered_teachers.append(teacher_map[code])
        return tuple(ordered_teachers)

    if schedule_teachers:
        return schedule_teachers
    return tuple(teacher_map.values())


def load_schedule_teacher_codes(folder: Path, teachers: tuple[TeacherRecord, ...]) -> tuple[str, ...]:
    found_codes: list[str] = []
    seen_codes: set[str] = set()

    for workbook_path in sorted(folder.glob("*.xlsx")):
        if workbook_path.name.startswith("~$"):
            continue
        workbook = load_workbook(workbook_path, data_only=True)
        if "座位区块" in workbook.sheetnames:
            continue
        for sheet_name in workbook.sheetnames:
            if sheet_name == "去重":
                continue
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows(values_only=True):
                for cell_value in row:
                    code = normalize_initial(cell_value)
                    if is_candidate_schedule_code(code) and code not in seen_codes:
                        found_codes.append(code)
                        seen_codes.add(code)
    return tuple(found_codes)


def resolve_teacher_records_from_codes(
    codes: tuple[str, ...],
    teachers: tuple[TeacherRecord, ...],
    include_unresolved: bool = True,
) -> tuple[TeacherRecord, ...]:
    teacher_map = build_teacher_lookup(teachers)
    resolved: list[TeacherRecord] = []
    seen_primary_codes: set[str] = set()

    for code in codes:
        teacher = teacher_map.get(code)
        if teacher is None:
            if not include_unresolved:
                continue
            resolved.append(
                TeacherRecord(
                    teacher_id="",
                    name="",
                    initial=code,
                    seat_no="",
                    group="",
                    unresolved=True,
                )
            )
            seen_primary_codes.add(code)
            continue

        primary_code = normalize_initial(teacher.initial)
        if primary_code not in seen_primary_codes:
            resolved.append(teacher)
            seen_primary_codes.add(primary_code)

    return tuple(resolved)


def is_image_input_path(input_path: Path) -> bool:
    return input_path.suffix.lower() in {".tif", ".tiff", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}


def detect_input_pdf_codes(
    input_pdf: Path,
    teachers: tuple[TeacherRecord, ...],
    config: AppConfig,
) -> tuple[str, ...]:
    teacher_map = build_teacher_lookup(teachers)
    found_codes: list[str] = []
    seen_codes: set[str] = set()
    ocr_engine: object | None = None

    with pymupdf.open(input_pdf) as document:
        for page in document:
            page_codes = detect_pdf_invigilator_codes_from_text_layer(page, teacher_map)
            if page_codes:
                for code in page_codes:
                    if code not in seen_codes:
                        found_codes.append(code)
                        seen_codes.add(code)
                continue

            if config.text_only_mode and document.is_pdf:
                continue

            if ocr_engine is None:
                ocr_engine = create_ocr_engine()
            for code in detect_pdf_invigilator_codes_from_ocr(page, ocr_engine, config, teacher_map):
                if code not in seen_codes:
                    found_codes.append(code)
                    seen_codes.add(code)

    return tuple(found_codes)


def detect_page_matches(
    page,
    ocr_engine,
    config: AppConfig,
    teacher_map: dict[str, TeacherRecord],
    allowed_teacher_codes: tuple[str, ...] | None = None,
):
    matches: dict[str, list[MatchOccurrence]] = defaultdict(list)
    if allowed_teacher_codes:
        teacher_map = {code: teacher_map[code] for code in allowed_teacher_codes if code in teacher_map}
    if not teacher_map:
        return matches, False, False

    text_matches = detect_text_layer_matches(page, teacher_map)
    if text_matches:
        return text_matches, True, False

    if ocr_engine is None:
        return matches, False, True

    segment_codes = build_segmentation_codes(teacher_map)
    detections = run_ocr_detections(page, ocr_engine, config)
    effective_threshold = get_effective_ocr_confidence_threshold(config)

    for rect, recognized_text, confidence in detections:
        if confidence < effective_threshold:
            continue

        segments = extract_candidate_codes(str(recognized_text).strip(), segment_codes)
        if not segments:
            continue

        for teacher_code, segment_rect in split_detection_rect(rect, segments):
            if teacher_code not in teacher_map:
                continue
            canonical_teacher = teacher_map[teacher_code]
            canonical_code = normalize_initial(canonical_teacher.initial)
            matches[canonical_code].append(
                MatchOccurrence(
                    page_index=page.number,
                    rect=segment_rect,
                    recognized_text=str(recognized_text).strip(),
                    confidence=confidence,
                    detection_source="ocr",
                )
            )
    return matches, False, False


def create_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def detect_text_layer_matches(page, teacher_map: dict[str, TeacherRecord]) -> dict[str, list[MatchOccurrence]]:
    words = page.get_text("words")
    if not words:
        return {}

    ordered_codes = build_segmentation_codes(teacher_map)
    matches: dict[str, list[MatchOccurrence]] = defaultdict(list)
    for word in words:
        rect = (float(word[0]), float(word[1]), float(word[2]), float(word[3]))
        text = str(word[4]).strip()
        if not text:
            continue
        segments = segment_teacher_text(text, ordered_codes)
        if not segments:
            continue
        for teacher_code, segment_rect in split_detection_rect(rect, segments):
            if teacher_code not in teacher_map:
                continue
            canonical_teacher = teacher_map[teacher_code]
            canonical_code = normalize_initial(canonical_teacher.initial)
            matches[canonical_code].append(
                MatchOccurrence(
                    page_index=page.number,
                    rect=segment_rect,
                    recognized_text=text,
                    confidence=1.0,
                    detection_source="text",
                )
            )
    return matches


def detect_codes_from_text_layer(page, ordered_codes: tuple[str, ...]) -> tuple[str, ...]:
    items = [
        ((float(word[0]), float(word[1]), float(word[2]), float(word[3])), str(word[4]).strip())
        for word in page.get_text("words")
        if str(word[4]).strip()
    ]
    return collect_codes_from_row_items(items, ordered_codes)


def detect_codes_from_ocr(page, ocr_engine, config: AppConfig, ordered_codes: tuple[str, ...]) -> tuple[str, ...]:
    effective_threshold = get_effective_ocr_confidence_threshold(config)
    items = [
        (rect, str(recognized_text).strip())
        for rect, recognized_text, confidence in run_ocr_detections(page, ocr_engine, config)
        if confidence >= effective_threshold and str(recognized_text).strip()
    ]
    return collect_codes_from_row_items(items, ordered_codes)


def detect_pdf_invigilator_codes_from_text_layer(
    page,
    teacher_map: dict[str, TeacherRecord],
) -> tuple[str, ...]:
    items = [
        ((float(word[0]), float(word[1]), float(word[2]), float(word[3])), str(word[4]).strip())
        for word in page.get_text("words")
        if str(word[4]).strip()
    ]
    return collect_pdf_invigilator_codes(items, teacher_map, strict_text_layer=True)


def detect_pdf_invigilator_codes_from_ocr(
    page,
    ocr_engine,
    config: AppConfig,
    teacher_map: dict[str, TeacherRecord],
) -> tuple[str, ...]:
    effective_threshold = get_effective_ocr_confidence_threshold(config)
    items = [
        (rect, str(recognized_text).strip())
        for rect, recognized_text, confidence in run_ocr_detections(page, ocr_engine, config)
        if confidence >= effective_threshold and str(recognized_text).strip()
    ]
    return collect_pdf_invigilator_codes(items, teacher_map, strict_text_layer=False)


def extract_candidate_codes(text: str, ordered_codes: tuple[str, ...]) -> tuple[str, ...]:
    normalized_text = normalize_initial(text)
    if not normalized_text:
        return ()

    segmented = segment_teacher_text(normalized_text, ordered_codes)
    if segmented:
        return segmented

    if is_candidate_schedule_code(normalized_text):
        return (normalized_text,)

    parts = re.split(r"[^A-Z\u4E00-\u9FFF]+", normalized_text)
    found: list[str] = []
    for part in parts:
        if not part:
            continue
        segmented = segment_teacher_text(part, ordered_codes)
        if segmented:
            found.extend(segmented)
        elif is_candidate_schedule_code(part):
            found.append(part)
    return tuple(found)


def collect_codes_from_row_items(
    items: list[tuple[tuple[float, float, float, float], str]],
    ordered_codes: tuple[str, ...],
) -> tuple[str, ...]:
    rows = build_text_rows(items)
    row_infos = []
    for row in rows:
        words = [text for _, text in row]
        candidate_codes: list[str] = []
        candidate_word_count = 0
        meaningful_word_count = 0
        for _, text in row:
            text = text.strip()
            if not text:
                continue
            meaningful_word_count += 1
            codes = extract_candidate_codes(text, ordered_codes)
            if codes:
                candidate_word_count += 1
                candidate_codes.extend(codes)
        row_infos.append(
            {
                "y": row[0][0][1],
                "candidate_codes": tuple(candidate_codes),
                "candidate_word_count": candidate_word_count,
                "meaningful_word_count": meaningful_word_count,
            }
        )

    qualifying_indexes: set[int] = set()
    for index, info in enumerate(row_infos):
        candidate_word_count = info["candidate_word_count"]
        meaningful_word_count = info["meaningful_word_count"]
        if candidate_word_count >= 2 and candidate_word_count >= max(2, int(meaningful_word_count * 0.5)):
            qualifying_indexes.add(index)

    for index, info in enumerate(row_infos):
        if index in qualifying_indexes:
            continue
        candidate_word_count = info["candidate_word_count"]
        meaningful_word_count = info["meaningful_word_count"]
        if candidate_word_count == 0 or candidate_word_count != meaningful_word_count:
            continue
        if has_nearby_candidate_row(row_infos, index):
            qualifying_indexes.add(index)

    found: list[str] = []
    seen: set[str] = set()
    for index in sorted(qualifying_indexes):
        for code in row_infos[index]["candidate_codes"]:
            if not strict_text_layer and code not in teacher_map:
                continue
            if code not in seen:
                found.append(code)
                seen.add(code)
    return tuple(found)


def collect_pdf_invigilator_codes(
    items: list[tuple[tuple[float, float, float, float], str]],
    teacher_map: dict[str, TeacherRecord],
    strict_text_layer: bool,
) -> tuple[str, ...]:
    ordered_codes = build_segmentation_codes(teacher_map)
    row_tolerance = TEXT_ROW_Y_TOLERANCE if strict_text_layer else OCR_ROW_Y_TOLERANCE
    nearby_row_tolerance = TEXT_NEARBY_ROW_TOLERANCE if strict_text_layer else OCR_NEARBY_ROW_TOLERANCE
    rows = build_text_rows(items, row_tolerance=row_tolerance)
    row_infos = []

    for row in rows:
        candidate_codes: list[str] = []
        candidate_word_count = 0
        meaningful_word_count = 0
        known_code_count = 0
        for _, text in row:
            text = text.strip()
            if not text:
                continue
            meaningful_word_count += 1
            codes = extract_pdf_check_candidate_codes(text, ordered_codes, strict_text_layer)
            if not codes:
                continue
            candidate_word_count += 1
            candidate_codes.extend(codes)
            known_code_count += sum(1 for code in codes if code in teacher_map)
        row_infos.append(
            {
                "y": row[0][0][1],
                "candidate_codes": tuple(candidate_codes),
                "candidate_word_count": candidate_word_count,
                "meaningful_word_count": meaningful_word_count,
                "known_code_count": known_code_count,
            }
        )

    qualifying_indexes: set[int] = set()
    for index, info in enumerate(row_infos):
        known_code_count = info["known_code_count"]
        candidate_word_count = info["candidate_word_count"]
        meaningful_word_count = info["meaningful_word_count"]
        if known_code_count >= 2:
            qualifying_indexes.add(index)
            continue
        if known_code_count >= 1 and candidate_word_count == meaningful_word_count:
            qualifying_indexes.add(index)
            continue
        if known_code_count >= 1 and candidate_word_count >= 2 and candidate_word_count >= max(2, int(meaningful_word_count * 0.5)):
            qualifying_indexes.add(index)
            continue
        if not strict_text_layer and known_code_count >= 1 and has_nearby_candidate_row(
            row_infos,
            index,
            max_distance=nearby_row_tolerance,
        ):
            qualifying_indexes.add(index)

    for index, info in enumerate(row_infos):
        if index in qualifying_indexes:
            continue
        candidate_word_count = info["candidate_word_count"]
        meaningful_word_count = info["meaningful_word_count"]
        if candidate_word_count == 0 or candidate_word_count != meaningful_word_count:
            continue
        if has_nearby_known_invigilator_row(row_infos, index, max_distance=nearby_row_tolerance):
            qualifying_indexes.add(index)

    found: list[str] = []
    seen: set[str] = set()
    for index in sorted(qualifying_indexes):
        for code in row_infos[index]["candidate_codes"]:
            if code not in seen:
                found.append(code)
                seen.add(code)
    return tuple(found)


def extract_pdf_check_candidate_codes(
    text: str,
    ordered_codes: tuple[str, ...],
    strict_text_layer: bool,
) -> tuple[str, ...]:
    normalized_text = normalize_initial(text)
    if not normalized_text:
        return ()

    segmented = segment_teacher_text(normalized_text, ordered_codes)
    if segmented:
        return segmented

    if strict_text_layer:
        if any(ch.isdigit() for ch in normalized_text):
            return ()
        if any(ch in normalized_text for ch in "()（）,，/:："):
            return ()
        if is_candidate_schedule_code(normalized_text):
            return (normalized_text,)
        return ()

    return extract_candidate_codes(text, ordered_codes)


def get_effective_ocr_confidence_threshold(config: AppConfig) -> float:
    return min(config.confidence_threshold, OCR_EFFECTIVE_CONFIDENCE_CAP)


def build_text_rows(
    items: list[tuple[tuple[float, float, float, float], str]],
    row_tolerance: float = TEXT_ROW_Y_TOLERANCE,
) -> list[list[tuple[tuple[float, float, float, float], str]]]:
    rows: list[list[tuple[tuple[float, float, float, float], str]]] = []
    for item in sorted(items, key=lambda value: (value[0][1], value[0][0])):
        if not rows or abs(item[0][1] - rows[-1][0][0][1]) > row_tolerance:
            rows.append([item])
        else:
            rows[-1].append(item)
    return rows


def has_nearby_candidate_row(row_infos: list[dict], index: int, max_distance: float = TEXT_NEARBY_ROW_TOLERANCE) -> bool:
    current_y = row_infos[index]["y"]
    for other_index in (index - 1, index + 1):
        if other_index < 0 or other_index >= len(row_infos):
            continue
        other = row_infos[other_index]
        if abs(other["y"] - current_y) <= max_distance and other["candidate_word_count"] >= 1:
            return True
    return False


def has_nearby_known_invigilator_row(row_infos: list[dict], index: int, max_distance: float = TEXT_NEARBY_ROW_TOLERANCE) -> bool:
    current_y = row_infos[index]["y"]
    for other_index in (index - 1, index + 1):
        if other_index < 0 or other_index >= len(row_infos):
            continue
        other = row_infos[other_index]
        if abs(other["y"] - current_y) <= max_distance and other["known_code_count"] >= 1:
            return True
    return False


def run_ocr_detections(page, ocr_engine, config: AppConfig) -> list[tuple[tuple[float, float, float, float], str, float]]:
    detections: list[tuple[tuple[float, float, float, float], str, float]] = []
    detections.extend(run_ocr_on_rect(page, page.rect, max(config.ocr_zoom, 1.0), ocr_engine))

    tile_zoom = max(config.ocr_zoom * TILE_ZOOM_MULTIPLIER, 4.5)
    for tile_rect in iter_page_tiles(page.rect):
        detections.extend(run_ocr_on_rect(page, tile_rect, tile_zoom, ocr_engine))
    return detections


def run_ocr_on_rect(page, clip_rect, zoom: float, ocr_engine) -> list[tuple[tuple[float, float, float, float], str, float]]:
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=clip_rect, alpha=False)
    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, pixmap.n)
    result, _ = ocr_engine(image)
    if not result:
        return []

    x_scale = clip_rect.width / pixmap.width
    y_scale = clip_rect.height / pixmap.height
    detections: list[tuple[tuple[float, float, float, float], str, float]] = []
    for box, recognized_text, confidence_text in result:
        x_values = [point[0] for point in box]
        y_values = [point[1] for point in box]
        rect = (
            min(x_values) * x_scale + clip_rect.x0,
            min(y_values) * y_scale + clip_rect.y0,
            max(x_values) * x_scale + clip_rect.x0,
            max(y_values) * y_scale + clip_rect.y0,
        )
        detections.append((rect, str(recognized_text).strip(), float(confidence_text)))
    return detections


def iter_page_tiles(page_rect):
    tile_width = page_rect.width / TILE_COLS
    tile_height = page_rect.height / TILE_ROWS
    for row in range(TILE_ROWS):
        for column in range(TILE_COLS):
            x0 = max(page_rect.x0, page_rect.x0 + tile_width * column - TILE_OVERLAP)
            x1 = min(page_rect.x1, page_rect.x0 + tile_width * (column + 1) + TILE_OVERLAP)
            y0 = max(page_rect.y0, page_rect.y0 + tile_height * row - TILE_OVERLAP)
            y1 = min(page_rect.y1, page_rect.y0 + tile_height * (row + 1) + TILE_OVERLAP)
            yield pymupdf.Rect(x0, y0, x1, y1)


def build_segmentation_codes(teacher_map: dict[str, TeacherRecord]) -> tuple[str, ...]:
    return tuple(sorted(teacher_map.keys(), key=lambda item: (-len(item), item)))


def segment_teacher_text(text: str, ordered_codes: tuple[str, ...]) -> tuple[str, ...]:
    normalized_text = normalize_initial(text)
    if not normalized_text:
        return ()

    code_set = frozenset(ordered_codes)
    max_code_length = max(len(item) for item in ordered_codes)

    @lru_cache(maxsize=None)
    def split(index: int) -> tuple[str, ...] | None:
        if index == len(normalized_text):
            return ()
        for end in range(min(len(normalized_text), index + max_code_length), index, -1):
            piece = normalized_text[index:end]
            if piece not in code_set:
                continue
            tail = split(end)
            if tail is not None:
                return (piece,) + tail
        return None

    segments = split(0)
    return segments or ()


def split_detection_rect(rect: tuple[float, float, float, float], segments: tuple[str, ...]):
    if not segments:
        return ()
    if len(segments) == 1:
        return ((segments[0], rect),)

    total_length = sum(len(item) for item in segments)
    x0, y0, x1, y1 = rect
    width = x1 - x0
    height = y1 - y0
    horizontal = width >= height

    segment_rects = []
    cursor = x0 if horizontal else y0
    axis_length = width if horizontal else height
    for index, token in enumerate(segments):
        ratio = len(token) / total_length
        next_cursor = cursor + axis_length * ratio
        if index == len(segments) - 1:
            next_cursor = x1 if horizontal else y1
        if horizontal:
            token_rect = (cursor, y0, next_cursor, y1)
        else:
            token_rect = (x0, cursor, x1, next_cursor)
        segment_rects.append((token, token_rect))
        cursor = next_cursor
    return tuple(segment_rects)


def build_teacher_matches(
    output_teachers: tuple[TeacherRecord, ...],
    occurrences_by_teacher: dict[str, list[MatchOccurrence]],
    seat_map: dict[str, str],
    config: AppConfig,
    warnings: list[str],
) -> list[TeacherMatch]:
    matches: list[TeacherMatch] = []
    for teacher in output_teachers:
        teacher_code = normalize_initial(teacher.initial)
        occurrences = occurrences_by_teacher.get(teacher_code, [])
        seat_no = normalize_seat_no(teacher.seat_no)
        block = seat_map.get(seat_no, "")
        if teacher.unresolved:
            warnings.append(f"代码 {teacher.initial} 在老师名单中找不到，请操作人员检查。")
        elif teacher.seat_no and not block:
            warnings.append(f"老师 {teacher.initial} {teacher.name} 的座号 {teacher.seat_no} 找不到对应区块。")
        if not occurrences:
            if teacher.unresolved:
                warnings.append(f"未知代码 {teacher.initial} 没有在 PDF 中识别到可标记位置。")
            else:
                warnings.append(f"老师 {teacher.initial} {teacher.name} 没有在 PDF 中识别到可标记位置。")
        deduped_occurrences = tuple(
            sorted(
                deduplicate_occurrences(occurrences),
                key=lambda item: (item.page_index, item.rect[1], item.rect[0]),
            )
        )
        matches.append(
            TeacherMatch(
                teacher=teacher,
                seat_block=block,
                occurrences=deduped_occurrences,
            )
        )

    return sorted(
        matches,
        key=lambda item: build_teacher_sort_key(item, config.group_order),
    )


def deduplicate_occurrences(occurrences: list[MatchOccurrence]) -> list[MatchOccurrence]:
    unique: list[MatchOccurrence] = []
    for item in sorted(occurrences, key=lambda value: (value.page_index, value.rect[1], value.rect[0])):
        if any(is_same_occurrence(item, existing) for existing in unique):
            continue
        unique.append(item)
    return unique


def is_same_occurrence(left: MatchOccurrence, right: MatchOccurrence) -> bool:
    if left.page_index != right.page_index:
        return False
    if normalize_initial(left.recognized_text) != normalize_initial(right.recognized_text):
        return False
    return (
        abs(left.rect[0] - right.rect[0]) < 4
        and abs(left.rect[1] - right.rect[1]) < 4
        and abs(left.rect[2] - right.rect[2]) < 4
        and abs(left.rect[3] - right.rect[3]) < 4
    )


def append_teacher_pages(output_doc, source_doc, teacher_match: TeacherMatch, config: AppConfig, font_path: str):
    occurrences_by_page: dict[int, list[MatchOccurrence]] = defaultdict(list)
    for occurrence in teacher_match.occurrences:
        occurrences_by_page[occurrence.page_index].append(occurrence)

    for page_index in resolve_target_page_indexes(teacher_match, source_doc.page_count):
        page = append_source_page(output_doc, source_doc, page_index)
        register_font(page, font_path)
        add_teacher_label(page, teacher_match, config, font_path)
        for occurrence in occurrences_by_page.get(page_index, []):
            add_occurrence_marker(page, occurrence, config)


def append_source_page(output_doc, source_doc, page_index: int):
    if source_doc.is_pdf:
        output_doc.insert_pdf(source_doc, from_page=page_index, to_page=page_index)
        return output_doc[-1]

    source_page = source_doc[page_index]
    output_page = output_doc.new_page(width=source_page.rect.width, height=source_page.rect.height)
    pixmap = source_page.get_pixmap(alpha=False)
    output_page.insert_image(output_page.rect, pixmap=pixmap)
    return output_page


def resolve_target_page_indexes(teacher_match: TeacherMatch, source_page_count: int) -> tuple[int, ...]:
    if teacher_match.page_indexes:
        if source_page_count == 1:
            return (0,)
        return teacher_match.page_indexes
    if source_page_count >= 1:
        return (0,)
    return ()


def register_font(page, font_path: str):
    if font_path:
        page.insert_font(fontname="marker_font", fontfile=font_path)


def add_teacher_label(page, teacher_match: TeacherMatch, config: AppConfig, font_path: str):
    label_height = config.label_height
    font_name = "marker_font" if font_path else "helv"
    gap = config.label_gap
    text_values = (
        build_label_name_value(teacher_match.teacher),
        build_label_code_value(teacher_match.teacher),
        build_label_seat_value(teacher_match.teacher),
    )
    text_widths = [
        estimate_label_text_width(text, config.label_font_size)
        for text in text_values
    ]
    box_width = max(sum(text_widths) + gap * (len(text_values) - 1) + 24, 220)
    if config.label_anchor == "bottom_left":
        label_left = config.label_x
        label_top = config.label_y
    else:
        label_left = config.label_x
        label_top = page.rect.y1 - config.label_y - label_height
    label_rect = pymupdf.Rect(
        label_left,
        label_top,
        label_left + box_width,
        label_top + label_height,
    )
    base_x = label_rect.x0
    baseline_y = label_rect.y0 + config.label_font_size + 4
    cursor_x = base_x + 12
    for text, width in zip(text_values, text_widths):
        page.insert_text(
            pymupdf.Point(cursor_x, baseline_y),
            text,
            fontsize=config.label_font_size,
            fontname=font_name,
            color=(0, 0, 0),
            overlay=True,
        )
        cursor_x += width + gap


def add_occurrence_marker(page, occurrence: MatchOccurrence, config: AppConfig):
    rect = pymupdf.Rect(*occurrence.rect)
    marker_rect = rect
    if occurrence.detection_source == "ocr":
        marker_rect = add_highlight_marker(page, rect, config)
    add_star_marker(page, marker_rect, config)


def add_highlight_marker(page, rect: pymupdf.Rect, config: AppConfig) -> pymupdf.Rect:
    padding = max(config.highlight_padding, 1.0)
    highlight_rect = pymupdf.Rect(
        rect.x0 - padding,
        rect.y0 - padding,
        rect.x1 + padding,
        rect.y1 + padding,
    ) & page.rect
    page.draw_rect(
        highlight_rect,
        color=config.highlight_border_rgb,
        fill=config.highlight_fill_rgb,
        width=1.1,
        overlay=True,
        stroke_opacity=0.9,
        fill_opacity=0.2,
    )
    return highlight_rect


def add_star_marker(page, marker_rect: pymupdf.Rect, config: AppConfig):
    rect = marker_rect
    marker_size = max(rect.height * config.star_scale, 11)
    center_y = (rect.y0 + rect.y1) / 2
    star_width = pymupdf.get_text_length("*", fontname="helv", fontsize=marker_size)
    marker_point = pymupdf.Point(
        rect.x0 - star_width - config.star_gap,
        center_y + marker_size * 0.35 + config.star_y_offset,
    )
    page.insert_text(
        marker_point,
        "*",
        fontsize=marker_size,
        fontname="helv",
        color=(0.85, 0.0, 0.0),
        overlay=True,
    )


def estimate_label_text_width(text: str, font_size: float) -> float:
    total = 0.0
    for char in text:
        total += estimate_character_width(char, font_size)
    return max(total, font_size)


def estimate_character_width(char: str, font_size: float) -> float:
    codepoint = ord(char)
    if char.isspace():
        return font_size * 0.35
    if 0x4E00 <= codepoint <= 0x9FFF:
        return font_size * 1.02
    if 0x3400 <= codepoint <= 0x4DBF:
        return font_size * 1.02
    if char.isdigit():
        return font_size * 0.58
    if "A" <= char <= "Z":
        return font_size * 0.68
    if "a" <= char <= "z":
        return font_size * 0.58
    return font_size * 0.5


def is_candidate_schedule_code(code: str) -> bool:
    if not code:
        return False
    if code in EXCLUDED_SCHEDULE_TOKENS:
        return False
    if any(ch.isdigit() for ch in code):
        return False
    if len(code) == 1 and is_cjk_character(code):
        return True
    if 2 <= len(code) <= 3 and all("A" <= ch <= "Z" for ch in code):
        return True
    return False


def is_cjk_character(char: str) -> bool:
    if len(char) != 1:
        return False
    codepoint = ord(char)
    return 0x3400 <= codepoint <= 0x9FFF


def build_label_name_value(teacher: TeacherRecord) -> str:
    if teacher.unresolved:
        return ""
    return teacher.name


def build_label_code_value(teacher: TeacherRecord) -> str:
    if teacher.unresolved:
        return f"!{teacher.initial}"
    return teacher.initial


def build_label_seat_value(teacher: TeacherRecord) -> str:
    if teacher.unresolved:
        return ""
    return teacher.seat_no


def build_teacher_sort_key(teacher_match: TeacherMatch, group_order: tuple[str, ...]):
    block = teacher_match.seat_block.upper()
    seat_no = normalize_seat_no(teacher_match.teacher.seat_no)
    if not block:
        return (99, 99, 99, seat_no, teacher_match.teacher.initial)

    match = re.fullmatch(r"(CL|CR|L|R)(\d+)", block)
    if not match:
        return (98, 98, 98, block, seat_no, teacher_match.teacher.initial)

    prefix, numeric_text = match.groups()
    index = int(numeric_text)
    family_name = "C" if prefix in {"CL", "CR"} else prefix
    family_position = {name.upper(): idx for idx, name in enumerate(group_order)}.get(family_name, 97)
    center_position = 0
    if prefix == "CR":
        center_position = 1
    return (family_position, index, center_position, seat_no, teacher_match.teacher.initial)


def normalize_initial(text: object) -> str:
    normalized = str(text or "").strip().upper().replace(" ", "")
    for source, target in CODE_NORMALIZATION_MAP.items():
        normalized = normalized.replace(source, target)
    return normalized


def build_teacher_lookup(teachers: tuple[TeacherRecord, ...]) -> dict[str, TeacherRecord]:
    lookup: dict[str, TeacherRecord] = {}
    alias_conflicts: set[str] = set()

    for teacher in teachers:
        lookup[normalize_initial(teacher.initial)] = teacher

    for teacher in teachers:
        for alias in generate_teacher_aliases(teacher):
            if alias in lookup and lookup[alias] is not teacher:
                alias_conflicts.add(alias)
                continue
            lookup[alias] = teacher

    for alias in alias_conflicts:
        primary_owner = next(
            (teacher for teacher in teachers if normalize_initial(teacher.initial) == alias),
            None,
        )
        if primary_owner is not None:
            lookup[alias] = primary_owner
        elif alias in lookup:
            del lookup[alias]

    return lookup


def generate_teacher_aliases(teacher: TeacherRecord) -> tuple[str, ...]:
    aliases: list[str] = []
    primary_code = normalize_initial(teacher.initial)
    if primary_code:
        aliases.append(primary_code)

    cjk_chars = [char for char in teacher.name.strip() if 0x3400 <= ord(char) <= 0x9FFF]
    if cjk_chars:
        aliases.append(normalize_initial(cjk_chars[-1]))

    return tuple(dict.fromkeys(alias for alias in aliases if alias))


def resolve_font_path(configured_path: str) -> str:
    if configured_path:
        path = Path(configured_path)
        if not path.exists():
            raise FileNotFoundError(f"设定的字型档不存在：{configured_path}")
        return str(path)

    for item in COMMON_FONT_CANDIDATES:
        if Path(item).exists():
            return item
    return ""
